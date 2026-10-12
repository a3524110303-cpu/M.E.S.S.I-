import io
from contextlib import ExitStack
from types import SimpleNamespace
from unittest.mock import patch

from django.core.exceptions import PermissionDenied
from django.test import TestCase, TransactionTestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from cloud import import_templates
from messi.data import ValidationError, load_csv, load_excel
from portal.tests.fixtures import SchoolFixture


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class GroupTemplateTests(SchoolFixture, TestCase):
    def rows(self):
        return import_templates.group_rows(self.teacher, self.group)

    def test_template_contains_only_students_enrolled_in_authorized_subject(self):
        rows = self.rows()
        self.assertEqual([row["id_estudiante"] for row in rows], [self.student.code])
        self.assertEqual(rows[0]["Estudiante"], self.student_user.get_full_name())
        with self.assertRaises(PermissionDenied):
            import_templates.group_rows(self.teacher, self.other_group)

    def test_excel_names_are_guidance_and_blank_students_are_not_imported(self):
        rows = self.rows() + [{"Estudiante": "Sin capturar", "id_estudiante": "EST-100"}]
        book = load_workbook(io.BytesIO(import_templates.excel_template(rows, self.group)))
        self.assertEqual(book["Instrucciones"]["B2"].value, self.group.subject)
        sheet = book["Datos"]
        self.assertEqual(sheet["B2"].value, self.student.code)
        self.assertIsNone(sheet["C2"].value)
        sheet["A2"] = "Nombre cambiado: no determina la identidad"
        sheet["C2"], sheet["D2"], sheet["E2"] = 7.5, 80, 90
        output = io.BytesIO()
        book.save(output)
        book.close()
        records = load_excel(output.getvalue(), allow_temporary_ids=False, allow_student_names=True, skip_blank_indicators=True)
        self.assertEqual(records, [{"id_estudiante": self.student.code, "nota_parcial": 7.5, "asistencia": 80.0, "tareas_entregadas": 90.0}])

    def test_student_name_that_looks_like_formula_is_exported_as_text(self):
        book = load_workbook(io.BytesIO(import_templates.excel_template([{"Estudiante": "=1+1", "id_estudiante": "EST-001"}], self.group)))
        self.assertEqual(book["Datos"]["A2"].data_type, "s")
        self.assertEqual(book["Datos"]["A2"].value, "=1+1")
        book.close()

    def test_csv_is_prefilled_with_codes_and_partial_capture_is_rejected(self):
        csv = import_templates.csv_template(self.rows())
        self.assertIn(self.student.code + ",,,", csv)
        with self.assertRaises(ValidationError):
            load_csv(csv, skip_blank_indicators=True)
        with self.assertRaisesRegex(ValidationError, "asistencia está vacío"):
            load_csv(csv.replace(self.student.code + ",,,", self.student.code + ",7,,"), skip_blank_indicators=True)

    def test_zero_indicators_are_valid_and_not_treated_as_blank(self):
        csv = import_templates.csv_template(self.rows()).replace(self.student.code + ",,,", self.student.code + ",0,0,0")
        self.assertEqual(len(load_csv(csv, skip_blank_indicators=True)), 1)

    def test_default_excel_contract_remains_strict(self):
        raw = import_templates.excel_template(self.rows(), self.group)
        with self.assertRaises(ValidationError):
            load_excel(raw)


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class ImportRosterUITests(SchoolFixture, TransactionTestCase):
    def setUp(self):
        type(self).setUpTestData()

    def app(self, upload_content=None):
        from pathlib import Path
        from streamlit.testing.v1 import AppTest

        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=20)
        app.session_state["_messi_auth_user_id"] = self.teacher.pk
        app.session_state["_messi_auth_hash"] = self.teacher.get_session_auth_hash()
        app.session_state["_messi_auth_started"] = timezone.now().timestamp()
        app.session_state["_messi_auth_role"] = "docente"
        with ExitStack() as patches:
            patches.enter_context(patch("cloud.runtime.initialize", return_value=True))
            if upload_content is not None:
                upload = SimpleNamespace(name="indicadores.csv", size=len(upload_content), getvalue=lambda: upload_content)
                patches.enter_context(patch("streamlit.file_uploader", return_value=upload))
                patches.enter_context(patch("streamlit.form_submit_button", side_effect=lambda label, **kwargs: label == "Revisar archivo"))
            app.run()
        self.assertEqual(len(app.exception), 0)
        return app

    def test_teacher_sees_subject_names_codes_and_ready_templates(self):
        app = self.app()
        self.assertIn("Importar el primer parcial de Matemáticas", [item.value for item in app.subheader])
        self.assertTrue(any("no a la materia" in item.value for item in app.info))
        roster = next(item.value for item in app.dataframe if "Código para la plantilla" in item.value.columns)
        self.assertEqual(roster["Código para la plantilla"].tolist(), [self.student.code])
        self.assertEqual(roster["Estudiante"].tolist(), [self.student_user.get_full_name()])
        self.assertNotIn(self.other_student.code, roster.to_string())

    def test_empty_group_explains_enrollment_and_does_not_offer_upload(self):
        self.indicator.delete()
        self.enrollment.delete()
        app = self.app()
        self.assertTrue(any("todavía no tiene alumnos inscritos" in item.value for item in app.warning))
        self.assertEqual(len(app.get("file_uploader")), 0)

    def test_wrong_codes_are_named_without_revealing_other_students(self):
        app = self.app(b"id_estudiante,nota_parcial,asistencia,tareas_entregadas\nEST-002,7,80,90\n")
        errors = "\n".join(item.value for item in app.error)
        self.assertIn("EST-002", errors)
        self.assertIn("Matemáticas, grupo 1A", errors)
        self.assertIn("pide al administrador que lo inscriba", errors)
        self.assertNotIn(self.other_student_user.get_full_name(), errors)
        self.assertNotIn(f"messi-ui-{self.teacher.pk}-import-{self.group.pk}-preview", app.session_state)

    def test_preview_uses_registered_name_and_friendly_indicator_labels(self):
        app = self.app(b"id_estudiante,nota_parcial,asistencia,tareas_entregadas\nEST-001,7,80,90\n")
        preview = next(item.value for item in app.dataframe if "Nota primer parcial" in item.value.columns)
        self.assertEqual(preview["Estudiante"].tolist(), [self.student_user.get_full_name()])
        self.assertEqual(preview["Código"].tolist(), [self.student.code])

    def test_empty_template_tells_teacher_which_values_to_fill(self):
        app = self.app(b"id_estudiante,nota_parcial,asistencia,tareas_entregadas\nEST-001,,,\n")
        self.assertTrue(any("Llena nota, asistencia y tareas" in item.value for item in app.error))

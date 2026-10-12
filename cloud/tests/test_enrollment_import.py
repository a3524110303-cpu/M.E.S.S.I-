import io
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.test import TestCase, TransactionTestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook

from cloud import enrollment_files
from portal import enrollment_import
from portal.models import AuditEvent, Enrollment, Profile, Student
from portal.tests.fixtures import SchoolFixture


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class EnrollmentImportTests(SchoolFixture, TestCase):
    def rows(self):
        return [{"codigo": "", "nombres": "Lucía", "apellidos": "Pérez", "usuario": ""},
                {"codigo": self.student.code, "nombres": "", "apellidos": "", "usuario": ""}]

    def preview(self, rows=None):
        return enrollment_import.prepare(self.admin, self.rows() if rows is None else rows, [self.group.pk, self.other_group.pk])

    def test_review_does_not_write_and_commit_creates_accounts_and_multiple_subjects(self):
        before = get_user_model().objects.count()
        preview = self.preview()
        self.assertEqual(get_user_model().objects.count(), before)
        self.assertEqual((preview["new_accounts"], preview["new_enrollments"], preview["skipped"]), (1, 3, 1))
        previous_password = self.student_user.password
        result = enrollment_import.commit(self.admin, preview)
        access = result["credentials"][0]
        user = get_user_model().objects.get(username=access["usuario"])
        self.assertTrue(user.check_password(access["clave"]))
        self.assertEqual(Profile.objects.get(user=user).role, "estudiante")
        self.assertFalse(user.is_staff or user.is_superuser)
        self.assertEqual(Enrollment.objects.filter(student=user.student).count(), 2)
        self.assertEqual(AuditEvent.objects.filter(action="inscribir_estudiante").count(), 3)
        self.student_user.refresh_from_db()
        self.assertEqual(self.student_user.password, previous_password)
        self.assertEqual(result["enrolled"], 3)

    def test_existing_students_can_be_reimported_without_duplicates_or_password_reset(self):
        rows = [{"codigo": self.student.code}]
        result = enrollment_import.commit(self.admin, self.preview(rows))
        self.assertEqual(result["enrolled"], 1)
        result = enrollment_import.commit(self.admin, self.preview(rows))
        self.assertEqual(result, {"credentials": [], "enrolled": 0, "skipped": 2})

    def test_errors_do_not_leave_partial_accounts_enrollments_or_audits(self):
        before = (get_user_model().objects.count(), Enrollment.objects.count(), AuditEvent.objects.count())
        preview = self.preview()
        with patch("portal.enrollment_import.AuditEvent.objects.bulk_create", side_effect=ValidationError("fallo")):
            with self.assertRaises(ValidationError):
                enrollment_import.commit(self.admin, preview)
        self.assertEqual(before, (get_user_model().objects.count(), Enrollment.objects.count(), AuditEvent.objects.count()))

    def test_invalid_conflicting_and_duplicate_rows_are_rejected(self):
        for rows in ([{"codigo": "MAL", "nombres": "A", "apellidos": "B"}],
                     [{"codigo": "EST-100"}],
                     [{"codigo": self.student.code, "nombres": "Otra persona"}],
                     [{"codigo": self.student.code}, {"codigo": self.student.code}],
                     [{"nombres": "A", "apellidos": "B", "usuario": self.teacher.username}],
                     [{"nombres": "A", "apellidos": "B", "usuario": "repetido"}, {"nombres": "C", "apellidos": "D", "usuario": "repetido"}]):
            with self.subTest(rows=rows), self.assertRaises(ValidationError):
                self.preview(rows)

    def test_only_admin_can_review_or_commit(self):
        preview = self.preview()
        for actor in (self.teacher, self.student_user, self.tutor):
            with self.assertRaises(PermissionDenied):
                enrollment_import.prepare(actor, self.rows(), [self.group.pk])
            with self.assertRaises(PermissionDenied):
                enrollment_import.commit(actor, preview)

    def test_closed_period_and_stale_preview_reject_all_writes(self):
        preview = self.preview()
        self.period.is_open = False
        self.period.save()
        with self.assertRaises(ValidationError):
            enrollment_import.commit(self.admin, preview)
        self.assertFalse(Student.objects.filter(code=preview["rows"][0]["codigo"]).exists())
        self.period.is_open = True
        self.period.save()
        Enrollment.objects.create(student=self.student, group=self.other_group)
        with self.assertRaisesRegex(ValidationError, "cambiaron"):
            enrollment_import.commit(self.admin, preview)

    def test_auto_codes_reserve_explicit_rows_and_new_account_collision_requires_review(self):
        rows = [{"nombres": "A", "apellidos": "B"}, {"codigo": "EST-003", "nombres": "C", "apellidos": "D"}]
        preview = self.preview(rows)
        self.assertEqual([row["codigo"] for row in preview["rows"]], ["EST-004", "EST-003"])
        get_user_model().objects.create_user(preview["rows"][0]["usuario"])
        with self.assertRaises(ValidationError):
            enrollment_import.commit(self.admin, preview)
        self.assertFalse(Student.objects.filter(code="EST-003").exists())

    def test_excel_and_csv_templates_roundtrip_accents_and_reject_formulas_and_extra_data(self):
        raw = enrollment_files.workbook(self.rows())
        self.assertEqual(enrollment_files.parse(raw, "lista.xlsx"), self.rows())
        csv = enrollment_files.csv_template() + ',Lucía,Pérez,\nEST-001,,,\n'
        self.assertEqual(enrollment_files.parse(csv.encode("utf-8"), "lista.csv"), self.rows())
        book = load_workbook(io.BytesIO(raw))
        book["Alumnos"]["B2"] = "=1+1"
        output = io.BytesIO()
        book.save(output)
        book.close()
        with self.assertRaisesRegex(ValidationError, "fórmulas"):
            enrollment_files.parse(output.getvalue(), "lista.xlsx")
        for content in (b"codigo,nombres,apellidos,usuario\n", b"codigo,nombres,apellidos,usuario\n,A,B,,extra\n", b"nombres,apellidos\nA,B\n"):
            with self.assertRaises(ValidationError):
                enrollment_files.parse(content, "lista.csv")

    def test_limits_and_credentials_export_as_text(self):
        with self.assertRaises(ValidationError):
            enrollment_files.parse(b"x" * (enrollment_files.MAX_BYTES + 1), "lista.csv")
        with self.assertRaises(ValidationError):
            self.preview([{"nombres": "A", "apellidos": "B"}] * 1001)
        book = load_workbook(io.BytesIO(enrollment_files.workbook([{"nombres": "=1+1", "clave": "secreto"}], credentials=True)))
        self.assertEqual(book["Accesos"]["B2"].data_type, "s")
        self.assertEqual(book["Accesos"]["E2"].value, "secreto")
        book.close()


@override_settings(PASSWORD_HASHERS=["django.contrib.auth.hashers.MD5PasswordHasher"])
class EnrollmentImportUITests(SchoolFixture, TransactionTestCase):
    def setUp(self):
        type(self).setUpTestData()

    def test_admin_upload_review_save_and_prevent_repeat(self):
        from streamlit.testing.v1 import AppTest
        app = AppTest.from_file(str(Path(__file__).resolve().parents[1] / "app.py"), default_timeout=20)
        app.session_state["_messi_auth_user_id"] = self.admin.pk
        app.session_state["_messi_auth_hash"] = self.admin.get_session_auth_hash()
        app.session_state["_messi_auth_started"] = timezone.now().timestamp()
        app.session_state["_messi_auth_role"] = "admin"
        raw = b"codigo,nombres,apellidos,usuario\n,Nuevo,Alumno,\n"
        uploaded = SimpleNamespace(name="lista.csv", size=len(raw), getvalue=lambda: raw)
        with patch("cloud.runtime.initialize", return_value=True), patch("streamlit.file_uploader", return_value=uploaded):
            app.run()
            self.assertEqual(len(app.exception), 0)
            app.multiselect(key=f"messi-ui-{self.admin.pk}-enrollment-file-groups").set_value([self.group.pk]).run()
            app.button(key=f"messi-ui-{self.admin.pk}-enrollment-file-review").click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any("1 cuentas nuevas" in item.value for item in app.success))
            app.button(key=f"messi-ui-{self.admin.pk}-enrollment-file-commit").click().run()
            self.assertEqual(len(app.exception), 0)
            self.assertTrue(any("Carga completada" in item.value for item in app.success))
            app.button(key=f"messi-ui-{self.admin.pk}-enrollment-file-review").click().run()
            self.assertTrue(any("ya se guardó" in item.value for item in app.error))

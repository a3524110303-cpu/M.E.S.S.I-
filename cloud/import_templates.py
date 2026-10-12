"""Plantillas de captura con los códigos y nombres del grupo autorizado."""
import csv
import io

from django.utils.text import slugify

FIELDS = ["id_estudiante", "nota_parcial", "asistencia", "tareas_entregadas"]
LIMIT = 10_000


def group_rows(user, group):
    from django.core.exceptions import ValidationError
    from portal import services

    enrollments = list(services.teacher_enrollments(user, group.pk).order_by("student__code")[:LIMIT + 1])
    if len(enrollments) > LIMIT:
        raise ValidationError("El grupo supera los 10,000 estudiantes. Solicita una plantilla por lotes al administrador.")
    return [{"id_estudiante": item.student.code,
             "Estudiante": item.student.user.get_full_name() or item.student.user.username}
            for item in enrollments]


def filename(group, suffix):
    label = slugify(f"{group.name}-{group.subject}-{group.period.name}")[:100]
    return f"plantilla_{label}.{suffix}"


def csv_template(rows):
    output = io.StringIO(newline="")
    output.write("\ufeff")
    writer = csv.writer(output)
    writer.writerow(FIELDS)
    writer.writerows([item["id_estudiante"], "", "", ""] for item in rows)
    return output.getvalue()


def excel_template(rows, group):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    book = Workbook()
    sheet = book.active
    sheet.title = "Datos"
    sheet.append(["Estudiante", "ID estudiante", "Nota primer parcial", "Asistencia (%)", "Tareas entregadas (%)"])
    for item in rows:
        sheet.append([item["Estudiante"], item["id_estudiante"], None, None, None])
        # Un nombre que empieza con '=' es texto, nunca una fórmula de Excel.
        sheet.cell(sheet.max_row, 1).data_type = "s"
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:E{len(rows) + 1}"
    for column, width in (("A", 36), ("B", 20), ("C", 23), ("D", 20), ("E", 25)):
        sheet.column_dimensions[column].width = width
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1B6751")
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for column, upper in (("C", 10), ("D", 100), ("E", 100)):
        rule = DataValidation(type="decimal", operator="between", formula1=0, formula2=upper, allow_blank=True)
        rule.showErrorMessage = True
        rule.errorTitle = "Valor fuera de rango"
        rule.error = f"Escribe un número entre 0 y {upper}."
        rule.showInputMessage = True
        rule.promptTitle = "Completa el indicador"
        rule.prompt = f"De 0 a {upper}." + (" Escribe 80 para 80 %." if upper == 100 else "")
        sheet.add_data_validation(rule)
        if rows:
            rule.add(f"{column}2:{column}{len(rows) + 1}")
    instructions = book.create_sheet("Instrucciones")
    for row in (
        ["Plantilla de MESSI", "Primer parcial"],
        ["Materia", group.subject], ["Grupo", group.name], ["Periodo", group.period.name],
        ["1", "En Datos, los nombres y códigos ya corresponden a los alumnos inscritos en esta materia."],
        ["2", "Conserva los códigos. Completa nota, asistencia y tareas de los alumnos que vas a importar."],
        ["3", "Las filas con los tres indicadores vacíos no se importan. Una fila parcialmente llena debe completarse."],
        ["4", "Usa nota de 0 a 10 y porcentajes de 0 a 100; escribe 80 para 80 %."],
        ["5", "Guarda como .xlsx, súbelo en la misma materia, revisa los nombres y confirma la importación."],
    ):
        instructions.append(row)
    instructions.column_dimensions["A"].width = 22
    instructions.column_dimensions["B"].width = 100
    for row in instructions:
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()

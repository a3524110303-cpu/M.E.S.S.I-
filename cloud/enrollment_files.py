"""Archivos de inscripción: una fila por alumno, sin contraseñas en la entrada."""
import csv
import io
from zipfile import ZipFile, BadZipFile

from django.core.exceptions import ValidationError

FIELDS = ("codigo", "nombres", "apellidos", "usuario")
LIMIT = 1000
MAX_BYTES = 5 * 1024 * 1024


def parse(content, filename):
    if len(content) > MAX_BYTES:
        raise ValidationError("El archivo debe pesar como máximo 5 MB.")
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            with ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * 1024 * 1024:
                    raise ValidationError("El Excel es demasiado grande al descomprimirlo.")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        except (BadZipFile, ValueError, OSError) as exc:
            raise ValidationError("No se pudo leer el Excel. Usa la plantilla .xlsx.") from exc
        try:
            sheet = book["Alumnos"] if "Alumnos" in book.sheetnames else book.active
            rows = []
            for cells in sheet.iter_rows():
                if len(rows) > LIMIT:
                    raise ValidationError(f"Usa como máximo {LIMIT} filas por archivo.")
                if any(cell.data_type == "f" for cell in cells):
                    raise ValidationError("El archivo contiene fórmulas. Pega los datos como valores.")
                rows.append([cell.value for cell in cells])
        finally:
            book.close()
    elif filename.lower().endswith(".csv"):
        try:
            decoded = content.decode("utf-8-sig")
            dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t")
            rows = []
            for row in csv.reader(io.StringIO(decoded), dialect):
                if len(rows) > LIMIT:
                    raise ValidationError(f"Usa como máximo {LIMIT} filas por archivo.")
                rows.append(row)
        except (UnicodeError, csv.Error) as exc:
            raise ValidationError("Guarda el archivo como CSV UTF-8 usando la plantilla.") from exc
    else:
        raise ValidationError("Sube un archivo .xlsx o .csv.")
    if not rows:
        raise ValidationError("El archivo está vacío.")
    headers = [str(value or "").strip().lower() for value in rows[0]]
    while headers and not headers[-1]:
        headers.pop()
    if len(headers) != len(set(headers)) or set(headers) != set(FIELDS):
        raise ValidationError("Usa las columnas codigo, nombres, apellidos y usuario de la plantilla.")
    records = []
    for number, row in enumerate(rows[1:], 2):
        if not any(value is not None and str(value).strip() for value in row):
            continue
        if any(value is not None and str(value).strip() for value in row[len(headers):]):
            raise ValidationError(f"Fila {number}: hay datos fuera de las columnas de la plantilla.")
        values = [str(value).strip() if value is not None else "" for value in row[:len(headers)]]
        values += [""] * (len(headers) - len(values))
        records.append(dict(zip(headers, values)))
    if not records:
        raise ValidationError("Completa al menos un alumno en la plantilla.")
    return records


def csv_template():
    return "\ufeff" + ",".join(FIELDS) + "\n"


def workbook(rows=None, credentials=False):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    book = Workbook()
    sheet = book.active
    sheet.title = "Accesos" if credentials else "Alumnos"
    fields = (*FIELDS, "clave") if credentials else FIELDS
    sheet.append(fields)
    for row in rows or []:
        sheet.append([row.get(field, "") for field in fields])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B6751")
    for column in ("A", "B", "C", "D", "E"):
        sheet.column_dimensions[column].width = 30
    sheet.freeze_panes = "A2"
    if not credentials:
        notes = book.create_sheet("Instrucciones")
        for line in (
            "Una fila por alumno. Todos se inscribirán en las materias seleccionadas en MESSI.",
            "Alumno nuevo: completa nombres y apellidos. Puedes dejar codigo y usuario vacíos: se generan automáticamente.",
            "Alumno existente: escribe su código EST-. Nombres, apellidos y usuario pueden quedar vacíos.",
            "Si completas datos de un alumno existente, deben coincidir con su cuenta; no se modificará.",
            "Conserva las cuatro columnas. Pega valores, sin fórmulas. Máximo 1000 alumnos por archivo.",
            "Sube el archivo y revisa antes de guardar. Descarga los accesos de las cuentas nuevas después de guardar.",
            "Para volver a importar alumnos creados automáticamente, usa los códigos del archivo de accesos.",
        ):
            notes.append([line])
        notes.column_dimensions["A"].width = 125
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()

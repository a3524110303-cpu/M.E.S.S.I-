"""Reconstruye los manuales de la versión Cloud vigente al 12 de octubre de 2026.

Ejecutar con Python y python-docx. No cambia código, datos ni credenciales.
"""
from pathlib import Path
import json

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TEAM = "Marco Antonio Osorio Hernandez; Ismael Hernández Jiménez; Víctor Manuel Jiménez Suárez; Yokio Yosafat Vazquez Carrillo; Salomón Alvarez Gomez"
URL = "https://messi-escolar-a3524110303.streamlit.app/"


def make_document(target, title, subtitle):
    # Reutiliza las dimensiones y estilos del manual del proyecto, reconstruyendo
    # el contenido por la petición de actualización completa.
    doc = Document(target)
    for element in list(doc._element.body):
        if element.tag != qn("w:sectPr"):
            doc._element.body.remove(element)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.8)
    section.left_margin = section.right_margin = Inches(.85)
    section.header_distance = section.footer_distance = Inches(.35)
    for name, size in (("Normal", 10.5), ("Title", 25), ("Subtitle", 11), ("Heading 1", 15), ("Heading 2", 11.5), ("Heading 3", 10.5)):
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.line_spacing = 1.12
        if name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.space_before = Pt(10)
    doc.styles["Normal"].paragraph_format.widow_control = True
    for hf in (section.header, section.footer):
        for paragraph in hf.paragraphs:
            paragraph.clear()
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("MESSI  |  ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = title + " de MESSI"
    doc.core_properties.subject = "Portal compartido y aplicación de escritorio"
    doc.core_properties.author = "Equipo MESSI"
    doc.core_properties.keywords = "MESSI, manual, Cloud, inscripción por archivo"
    doc.core_properties.comments = ""
    doc.add_paragraph(title + " de MESSI", "Title")
    doc.add_paragraph(subtitle, "Subtitle")
    p(doc, "Edición 2.0 · 12 de octubre de 2026 · Aplicación 0.4.0 con actualizaciones del portal Cloud")
    p(doc, "Equipo MESSI")
    p(doc, TEAM)
    return doc


def p(doc, text):
    return doc.add_paragraph(text)


def h(doc, text, level=1):
    return doc.add_paragraph(text, f"Heading {level}")


def page(doc, title):
    doc.add_page_break()
    h(doc, title)


def steps(doc, items):
    for number, text in enumerate(items, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(.2)
        paragraph.paragraph_format.first_line_indent = Inches(-.2)
        paragraph.add_run(f"{number}. ").bold = True
        paragraph.add_run(text)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    widths = widths or [6.8 / len(headers)] * len(headers)
    for column, width in zip(t.columns, widths):
        column.width = Inches(width)
    for row in [headers, *rows]:
        cells = t.rows[0].cells if row is headers else t.add_row().cells
        for cell, text, width in zip(cells, row, widths):
            cell.width = Inches(width)
            cell.text = str(text)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(5)
                paragraph.paragraph_format.space_before = Pt(4)
                paragraph.paragraph_format.line_spacing = 1.04
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = row is headers
            if row is headers:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "E8ECE9")
                cell._tc.get_or_add_tcPr().append(shading)
        props = cells[0]._tc.getparent().get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        if row is headers:
            props.append(OxmlElement("w:tblHeader"))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def code(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(9)
    for i, line in enumerate(text.splitlines()):
        run = paragraph.add_run(("\n" if i else "") + line)
        run.font.name = "Consolas"
        run.font.size = Pt(8)


def picture(doc, path, caption, width=6.1, crop=None):
    if not path.exists():
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    if crop:
        left, top, right, bottom = crop
        rect = OxmlElement("a:srcRect")
        for key, value in zip(("l", "t", "r", "b"), crop):
            rect.set(key, str(value))
        fill = shape._inline.xpath(".//pic:blipFill")[0]
        fill.insert(1, rect)
        shape.height = int(shape.height * (100000 - top - bottom) / (100000 - left - right))
    cap = doc.add_paragraph(caption)
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        run.font.size = Pt(9)


def user_manual():
    target = ROOT / "docs/entrega_3/04_Manual_de_usuario.docx"
    d = make_document(target, "Manual de usuario", "Guía de operación del portal escolar en línea")
    h(d, "1 Qué hace MESSI")
    p(d, "MESSI reúne indicadores del primer parcial, solicitudes de apoyo, acuerdos y seguimiento. Este manual permite preparar la escuela, inscribir una lista de alumnos y trabajar desde cada cuenta. El portal en línea comparte los registros entre dispositivos; cada persona ve únicamente el espacio y las asignaciones de su rol.")
    p(d, "La puntuación del modelo procede de datos sintéticos. Sirve para demostrar el recorrido de análisis y acompañamiento; no es una probabilidad validada de reprobación ni decide calificaciones, sanciones o apoyos automáticamente.")
    table(d, ["Rol", "Operaciones principales"], [
        ["Administrador", "Organiza cuentas, periodos, grupos, inscripciones y tutores."],
        ["Docente", "Captura e importa indicadores de sus grupos y calcula el análisis demostrativo."],
        ["Estudiante", "Consulta sus indicadores, envía solicitudes y revisa acuerdos y avances."],
        ["Tutor", "Revisa a sus estudiantes asignados, acuerda apoyos y registra seguimiento."],
    ], [1.35, 5.4])
    h(d, "Recorrido del manual", 2)
    p(d, "Acceso y preparación escolar; inscripción por Excel o CSV; captura docente; consulta del estudiante y acompañamiento del tutor; interpretación de resultados; preguntas frecuentes y solución de problemas. La última sección explica la modalidad local de Windows.")

    page(d, "2 Antes de empezar e iniciar sesión")
    p(d, "Para usar el portal necesitas conexión a internet, un navegador actualizado y una cuenta asignada por la escuela. No hace falta instalar MESSI en cada equipo. Abre " + URL)
    steps(d, ["Escribe tu Usuario y Contraseña y pulsa Entrar. El sistema abre el espacio correspondiente a tu rol.", "Si no tienes cuenta o no aparece un grupo, solicita al administrador la cuenta o la asignación que falta. Crear una cuenta docente por sí solo no asigna materias.", "Pulsa Actualizar mi espacio cuando quieras consultar cambios guardados por otra persona. Usa Cerrar sesión al terminar, especialmente en un equipo compartido."])
    p(d, "La sesión dura como máximo ocho horas. Puede terminar antes si cambian tu contraseña, tu rol o el estado activo de la cuenta. Si se bloquea el acceso tras varios intentos fallidos, respeta el tiempo indicado antes de volver a intentar.")
    h(d, "3 Preparar la escuela", 1)
    p(d, "El administrador entra al espacio ORGANIZACIÓN ESCOLAR. Primero crea las cuentas del personal y la estructura del periodo; después inscribe alumnos y asigna tutores.")
    steps(d, ["Cuentas: completa Usuario nuevo, Nombre, Apellidos, Rol de la cuenta nueva y la contraseña con su confirmación. El código EST- corresponde sólo a estudiantes. Para nuevos alumnos de una lista puedes omitir esta captura individual y usar Inscripciones.", "Periodos: completa nombre y fechas; activa Captura abierta y pulsa Crear periodo. La fecha final no debe ser anterior a la inicial.", "Grupos: en Crear grupo y materia, indica grupo, materia y periodo; selecciona los docentes y pulsa Crear grupo. Cada combinación de grupo, materia y periodo constituye una asignación distinta.", "Para un grupo ya creado usa Asignar docente a un grupo existente. Busca la cuenta docente, selecciona el grupo y pulsa Asignar docente. Se conservan los demás docentes asignados.", "Inscripciones: usa la carga de una lista explicada en las siguientes páginas o el formulario individual de la parte inferior.", "Tutores: busca y selecciona al estudiante inscrito, el tutor y su periodo. Pulsa Asignar tutor. La tutoría se asigna por estudiante y periodo; no se crea al inscribir una lista."])
    p(d, "Comprueba la confirmación y las tablas después de cada operación. Los selectores individuales muestran hasta 200 opciones; usa la búsqueda y la paginación cuando una persona no aparezca. Un periodo cerrado conserva la consulta, pero impide nuevas capturas y las altas por archivo.")

    page(d, "4 Completar la lista de inscripción")
    p(d, "En Administrador → Inscripciones, pulsa Descargar plantilla de inscripción Excel. Abre la hoja Alumnos. Conserva los cuatro encabezados y escribe una persona por fila; no escribas materias ni contraseñas en esta hoja.")
    table(d, ["Columna", "Alumno nuevo", "Alumno existente"], [
        ["codigo", "Vacía para generar un EST- automáticamente, o un código nuevo válido.", "Su código EST- registrado. Es obligatorio para reutilizar su cuenta."],
        ["nombres", "Nombre o nombres del alumno.", "Vacía o exactamente igual a su cuenta."],
        ["apellidos", "Apellido o apellidos del alumno.", "Vacía o exactamente igual a su cuenta."],
        ["usuario", "Vacía para generarlo, o un usuario nuevo válido y disponible.", "Vacía o exactamente igual al usuario registrado."],
    ], [1.0, 2.85, 2.9])
    h(d, "Ejemplo de tres alumnos", 2)
    table(d, ["codigo", "nombres", "apellidos", "usuario"], [
        ["", "Juan Carlos", "Pérez López", ""], ["", "María", "Hernández García", ""], ["EST-001", "", "", ""],
    ], [1.2, 1.7, 2.55, 1.3])
    p(d, "Las celdas vacías del ejemplo deben quedar realmente vacías: no escribas «vacío», guiones ni el texto de una explicación. Las dos primeras filas crearán cuentas. La tercera reutilizará EST-001 si existe y tiene una cuenta activa de estudiante.")
    h(d, "Reglas del archivo", 2)
    p(d, "Guarda el Excel como .xlsx o usa la plantilla CSV guardada en UTF-8. Pega valores, sin fórmulas. El archivo admite hasta 1000 alumnos y 5 MB; las filas totalmente vacías se omiten. No repitas un código ni un usuario. El código debe ser EST- seguido de 3 a 8 dígitos.")
    p(d, "Si vuelves a usar una lista de alumnos creados automáticamente, completa los códigos que recibiste en el archivo de accesos. Dejar codigo vacío vuelve a solicitar un alta nueva: MESSI no identifica a una persona sólo por su nombre.")
    p(d, "La hoja Instrucciones acompaña la plantilla. La importación no cambia datos personales ni contraseñas de cuentas existentes; corrige cualquier diferencia antes de guardar.")

    page(d, "5 Revisar y guardar la inscripción por archivo")
    steps(d, ["En Inscripciones, elige Periodo para la inscripción por archivo. Debe estar abierto y tener los grupos creados.", "Selecciona las Materias donde se inscribirá toda la lista. Puedes elegir hasta 30. Todos los alumnos del archivo se inscribirán en cada materia seleccionada; usa archivos separados si las listas son diferentes.", "Sube la Lista de alumnos (.xlsx o .csv) y pulsa Revisar lista de inscripción.", "Revisa la tabla de códigos, nombres, usuarios y tipo de cuenta. El resumen indica cuentas nuevas, inscripciones nuevas e inscripciones ya existentes.", "Si todo es correcto, pulsa Guardar cuentas e inscripciones y espera la confirmación. Si una fila falla o los datos cambiaron desde la revisión, se cancela el lote completo. Corrige y vuelve a revisar.", "Pulsa Descargar accesos de las cuentas nuevas antes de cerrar sesión o limpiar el resultado. Ese Excel contiene los códigos, usuarios y contraseñas generadas. Entrega a cada alumno sólo su propio acceso y conserva los códigos para otras materias."])
    p(d, "Las inscripciones que ya existen se omiten; las cuentas existentes no reciben una contraseña nueva. Cambiar el archivo o las materias obliga a una nueva revisión. Subir un archivo no lo guarda por sí solo.")
    evidence = ROOT / "docs/entrega_3/evidencias/inscripcion_por_archivo_2026-10-12.png"
    picture(d, evidence, "Pantalla publicada de inscripción por archivo en MESSI", width=5.4, crop=(33000, 46000, 5000, 0))

    page(d, "6 Capturar e importar indicadores como docente")
    p(d, "En ESPACIO DOCENTE selecciona Grupo y materia, comprobando su periodo. Sólo aparecen los grupos asignados a tu cuenta y sus estudiantes inscritos. Si faltan alumnos, el administrador debe inscribirlos antes de cargar indicadores.")
    h(d, "Captura y análisis", 2)
    steps(d, ["Selecciona Estudiante de esta página. Comprueba su nombre y código.", "Completa Calificación de 0 a 10, Asistencia (%) de 0 a 100 y Tareas entregadas (%) de 0 a 100. Escribe 80 para 80 por ciento, no 0.80.", "Pulsa Guardar indicadores. El cálculo utiliza los datos guardados; editar el formulario sin guardarlo no actualiza el análisis.", "Pulsa Calcular riesgo demostrativo para ese estudiante y revisa el mensaje de confirmación y la tabla. Si cambias sus indicadores, vuelve a calcular."])
    h(d, "Importar archivo del primer parcial", 2)
    steps(d, ["Abre Importar archivo y descarga Descargar Excel con mis alumnos. El Excel ya contiene los nombres y los códigos de la materia seleccionada.", "En Datos completa Nota primer parcial, Asistencia (%) y Tareas entregadas (%). Conserva los códigos. No uses la plantilla de inscripción: es una operación distinta.", "Puedes dejar los tres indicadores vacíos para omitir un alumno. Si llenas sólo uno o dos, completa la fila; los valores cero sí se importan.", "Guarda como .xlsx, súbelo en la misma materia y pulsa Revisar archivo. También puedes usar el CSV con los códigos y encabezados de su plantilla.", "Revisa los nombres y valores y pulsa Guardar importación. Esta operación actualiza los indicadores de todos los alumnos incluidos y no calcula automáticamente sus puntuaciones."])
    table(d, ["ID estudiante", "Nota primer parcial", "Asistencia (%)", "Tareas entregadas (%)"], [["EST-001", "7.5", "80", "90"]], [1.5, 1.7, 1.6, 1.95])
    p(d, "Si alguien modificó los indicadores después de tu revisión, MESSI rechaza el lote para evitar sobrescribirlos. Descarta la revisión o recarga los datos y revisa de nuevo. El periodo debe estar abierto para guardar.")
    h(d, "Exportar indicadores", 2)
    p(d, "Abre Exportar indicadores, pulsa Preparar CSV y descarga el archivo ofrecido. Es una exportación de indicadores del grupo, no un respaldo completo de cuentas, solicitudes o seguimientos.")

    page(d, "7 Consultar y acompañar al estudiante")
    h(d, "Cuenta de estudiante", 2)
    p(d, "En ESPACIO ESTUDIANTE, Mis indicadores muestra tus materias y datos registrados. Mi acompañamiento presenta los acuerdos y los avances visibles. No tienes que escribir el código de otro alumno para consultar información; la cuenta determina tu identidad.")
    steps(d, ["En Mis solicitudes, selecciona un periodo abierto en el que estés inscrito.", "Escribe un mensaje que explique el apoyo que necesitas y pulsa Enviar solicitud. Comprueba que aparece en Mis solicitudes enviadas.", "Después consulta Mi acompañamiento y pulsa Actualizar mi espacio para ver los acuerdos y avances que el tutor haya guardado."])
    p(d, "Puedes solicitar apoyo sin una alerta del modelo. La solicitud no crea un caso automáticamente: el tutor debe revisarla y acordar el acompañamiento. Las notas internas del tutor no aparecen en tu cuenta.")
    h(d, "Cuenta de tutor", 2)
    p(d, "En ESPACIO TUTOR sólo se muestran los estudiantes asignados a tu cuenta en sus periodos. Revisa Solicitudes y Alertas; una señal vigente del modelo debe interpretarse junto con la conversación y el contexto del alumno.")
    steps(d, ["Abre Acordar apoyo. Elige el origen disponible: una solicitud, una alerta o una conversación. Selecciona al alumno y el periodo que correspondan.", "Escribe el Acuerdo de apoyo visible al estudiante. Registra las notas internas separadamente y pulsa Crear caso de apoyo.", "En Seguimiento selecciona el caso. Anota el avance y el próximo paso visibles al estudiante, las notas internas que procedan y el estado del caso.", "Pulsa Guardar seguimiento. Comprueba la confirmación y el historial; el estudiante podrá ver el acuerdo y el avance después de actualizar su espacio."])
    p(d, "Los estados del caso son Pendiente, En seguimiento y Cerrado. Si otro tutor actualizó el caso mientras tenías el formulario abierto, recarga antes de intentar guardar. Reasignar el tutor por estudiante y periodo transfiere sus casos al nuevo tutor y retira el acceso del anterior.")

    page(d, "8 Interpretar resultados y resolver dudas")
    p(d, "La tabla docente puede mostrar Sin indicadores, Pendiente de cálculo, Desactualizado, Revisar apoyo o Sin alerta del modelo. Una predicción queda desactualizada cuando cambia la versión de sus indicadores. Guarda los datos correctos y calcula de nuevo para obtener un análisis vigente.")
    p(d, "La puntuación interna va de 0 a 1 y la interfaz Cloud la presenta sobre 100. Por ejemplo, 0.58 equivale a 58 / 100; no significa una probabilidad real de reprobación. El umbral de la demostración es 0.5, equivalente a 50 / 100. Una alerta pide revisar el contexto; su ausencia no demuestra que el alumno no necesite apoyo.")
    table(d, ["Pregunta", "Respuesta"], [
        ["¿Necesito instalar el programa?", "Para el portal en línea, no. Necesitas internet y una cuenta escolar."],
        ["¿Se conserva la información al cerrar?", "Los registros confirmados quedan en la base compartida. Un formulario sin guardar puede perderse."],
        ["¿Puedo cargar alumnos e indicadores en el mismo archivo?", "No. La inscripción crea cuentas y relaciones; la plantilla docente carga nota, asistencia y tareas."],
        ["¿La carga masiva asigna tutor?", "No. Después de inscribir, el administrador asigna al tutor en Tutores."],
        ["¿Me llegarán avisos por correo?", "Este recorrido se consulta dentro del portal. Actualiza tu espacio; no depende de un correo automático."],
    ], [2.4, 4.35])
    h(d, "Solución de problemas", 2)
    table(d, ["Situación", "Qué hacer"], [
        ["El enlace muestra reposo", "Usa la opción de Streamlit para reactivar la aplicación y espera a que aparezca Iniciar sesión."],
        ["No hay grupo o alumno", "Revisa rol, docente asignado, inscripción, periodo y filtro de búsqueda con el administrador."],
        ["El archivo se rechaza", "Lee la fila o el código indicado; conserva encabezados, pega valores y corrige datos o duplicados."],
        ["Usuario existente en un alta", "Si es un alumno registrado, usa su código EST-. Si es otra persona, elige un usuario disponible."],
        ["El código no está en la materia", "Comprueba la plantilla docente o pide la inscripción en ese grupo y materia."],
        ["La operación no se completa", "No repitas altas a ciegas. Revisa la confirmación, actualiza y comunica al administrador el mensaje, la operación y la hora, sin enviar contraseñas."],
    ], [2.2, 4.55])

    page(d, "9 Usar la modalidad de escritorio en Windows")
    p(d, "El repositorio conserva el instalador MESSI-Setup-Windows-x64.exe como demostración local. Esta modalidad abre una ventana de control y un navegador en la misma computadora; utiliza SQLite y no se sincroniza con el portal Cloud. Sus vistas Docente, Estudiante y Tutor no constituyen cuentas autenticadas.")
    h(d, "Instalar y abrir", 2)
    steps(d, ["En un equipo Windows compatible de 64 bits, abre el instalador y sigue Siguiente, Instalar y Finalizar. Conserva la carpeta propuesta y el acceso directo si lo necesitas.", "Abre MESSI desde el escritorio o Inicio. Espera Todo listo y mantén abierta la ventana de control. Si cerraste la pestaña, pulsa Abrir MESSI.", "En Docente usa captura directa, Excel o CSV, tabla pegada o el ejemplo sintético. Revisa, pulsa Guardar indicadores, calcula el riesgo de demostración y descarga el reporte.", "En Estudiante escribe el código ficticio y envía la solicitud. En Tutor consulta esa solicitud, registra el acuerdo y guarda el seguimiento del apoyo.", "Antes de terminar usa Guardar respaldo y después Cerrar MESSI. Cerrar sólo la pestaña no detiene el servidor local."])
    h(d, "Datos locales y recuperación", 2)
    p(d, "Los datos de cada usuario de Windows se guardan en %LOCALAPPDATA%\\MESSI\\data\\messi.sqlite3. No se comparten automáticamente entre equipos. Guardar diagnóstico produce un reporte técnico; no reemplaza el respaldo.")
    p(d, "Para restaurar una copia local, cierra MESSI, conserva la base actual y aparta con ella los archivos messi.sqlite3-wal y messi.sqlite3-shm si existen. Coloca el respaldo como messi.sqlite3 y vuelve a abrir para comprobar los registros. Solicita ayuda al equipo si no reconoces estos archivos.")
    h(d, "Alcance y ayuda", 2)
    p(d, "Usa información ficticia para las demostraciones. La versión Windows y el portal tienen rutas y permisos distintos; una contraseña del portal no se utiliza en el selector local de vistas. El mantenimiento del hosting, de la base remota y de las cuentas corresponde al administrador técnico.")
    p(d, "Referencia del proyecto: https://github.com/a3524110303-cpu/M.E.S.S.I- . La edición de este manual corresponde al portal publicado y a la carga por archivo incorporada el 12 de octubre de 2026.")
    d.save(target)
    return target


def programmer_manual():
    target = ROOT / "docs/entrega_2/02_Manual_del_programador.docx"
    d = make_document(target, "Manual del programador", "Arquitectura desarrollo y mantenimiento de MESSI")
    p(d, "Este manual describe la implementación existente, sus contratos de datos y el procedimiento para modificarla y verificarla. La ruta principal es Streamlit Community Cloud con modelos y servicios Django sobre MySQL externo. El repositorio también conserva un portal Django autohospedado y la demostración local de Windows.")
    h(d, "1 Arquitectura general")
    table(d, ["Modalidad", "Entrada", "Persistencia y alcance"], [
        ["Cloud publicada", "cloud/app.py", "ORM y servicios de web/portal sobre MySQL externo. Roles y asignaciones individuales."],
        ["Django autohospedado", "web/manage.py y web/config", "Vistas, sesiones Django, administración, servidor WSGI y MySQL."],
        ["Escritorio Windows", "messi_desktop.py y app.py", "Servidor local y SQLite por usuario. Vistas de demostración sin autenticación escolar."],
    ], [1.65, 1.9, 3.2])
    p(d, "El navegador Cloud envía acciones a Streamlit. cloud/runtime.py configura Django una vez por proceso, comprueba la conexión y las migraciones y prepara el modelo demostrativo si falta. La interfaz llama a portal.organization, portal.enrollment_import y portal.services; estas capas autorizan y ejecutan las transacciones. El ORM persiste los cambios en MySQL. No se inicia un servidor Django dentro de Streamlit.")
    p(d, "Los registros escolares no viven en st.session_state: allí se guardan la identidad de sesión, filtros, formularios, vistas previas y resultados temporales. La inscripción produce cuentas y relaciones; la captura produce indicadores; la inferencia produce una predicción versionada; el tutor produce casos y seguimientos.")
    h(d, "Límites de la implementación", 2)
    p(d, "La solución actual no implementa los microservicios de una arquitectura futura ni valida el modelo con estudiantes reales. La aplicación conserva su versión de paquete 0.4.0; esta edición documenta las actualizaciones Cloud del 12 de octubre de 2026. No confundir versión de documento, versión de paquete y revisión Git.")

    page(d, "2 Módulos y funciones principales")
    table(d, ["Ruta", "Contrato principal"], [
        ["cloud/app.py", "main y vistas _admin, _teacher, _student, _tutor. Formularios por rol, selección y paginación."],
        ["cloud/runtime.py", "initialize y configuration_environment. Secrets, entorno Django, TLS, migraciones y modelo."],
        ["cloud/auth.py", "authenticate_credentials, get_session_user y logout. Credenciales, bloqueo y sesión Cloud."],
        ["cloud/presentation.py", "Cabeceras, nombres legibles y error. Mensajes públicos sin revelar trazas privadas."],
        ["cloud/enrollment_files.py", "parse, workbook y csv_template. Lectura y plantillas de inscripción; exportación de accesos nuevos."],
        ["cloud/import_templates.py", "group_rows, excel_template y csv_template. Plantillas docentes con alumnos autorizados."],
        ["web/portal/organization.py", "create_account, create_period, create_group, assign_teacher, enroll_student, assign_tutor."],
        ["web/portal/enrollment_import.py", "prepare y commit. Revisión de alumnos y guardado atómico de cuentas e inscripciones."],
        ["web/portal/services.py", "Consultas visible_*, save_indicator, import_indicators, calculate_risk, create_request, create_case y add_followup."],
        ["web/portal/models.py", "Entidades, validadores y restricciones del expediente escolar compartido."],
        ["web/portal/views.py y forms.py", "Interfaz Django autohospedada. Reutiliza los servicios y contratos del portal."],
        ["web/config/settings.py", "Entorno, base, seguridad, sesiones, estáticos y parámetros de bloqueo."],
        ["src/messi/data.py y model.py", "Validación de indicadores e inferencia del modelo local de demostración."],
        ["src/messi/sqlite_storage.py", "SQLiteStore del escritorio. No es la persistencia del portal Cloud."],
    ], [2.75, 4.0])
    p(d, "Rutas src/messi: desktop.py controla el servidor y la ventana; paths.py separa recursos de datos; diagnostics.py comprueba el paquete; windows_process.py supervisa el proceso. scripts/train_demo.py genera modelo y metadatos; scripts/build_windows.ps1 reconstruye el cliente Windows.")

    page(d, "3 Entidades y relaciones del portal")
    p(d, "Las tablas del portal pertenecen al ORM Django y a sus migraciones en web/portal/migrations. No equivalen al esquema SQLite del escritorio. Las claves foráneas de expedientes usan protección de borrado para conservar relaciones; las restricciones de unicidad evitan duplicar inscripciones o cortes.")
    table(d, ["Entidad", "Relación o restricción relevante"], [
        ["User y Profile", "Una cuenta Django y un perfil con rol admin, docente, estudiante o tutor."],
        ["Student", "Un estudiante por cuenta; code único con EST- y de 3 a 8 dígitos."],
        ["Period", "Nombre único, fechas válidas e is_open para habilitar escritura."],
        ["CourseGroup", "Grupo + materia + periodo únicos. Relación de varios docentes con varios grupos."],
        ["Enrollment", "Una relación única entre estudiante y grupo/materia."],
        ["TutorAssignment", "Un tutor por estudiante y periodo; exige una inscripción en ese periodo."],
        ["Indicator", "Inscripción + corte únicos; grade, attendance, homework, version y autor de cambio."],
        ["Prediction", "Indicador, versión de entrada, score, threshold, model_version y marca alert."],
        ["HelpRequest", "Estudiante, periodo, mensaje y estado Pendiente o Atendida."],
        ["SupportCase", "Estudiante, periodo, tutor, acuerdo, notas internas, origen opcional y versión."],
        ["FollowUp", "Caso, autor, avance visible, nota interna, estado y fecha."],
        ["AuditEvent", "Actor, acción, entidad, identificador y fecha del cambio."],
        ["LoginAttempt", "Contador y bloqueo por clave normalizada de cuenta."],
    ], [1.75, 5.0])
    p(d, "Prediction.is_current compara indicator_version con Indicator.version. Una edición conserva el historial de predicciones pero deja obsoletas las anteriores; la vista debe distinguir ese estado. El comportamiento SQLite local puede eliminar la predicción anterior y debe documentarse por separado.")
    p(d, "El código EST- identifica al alumno, no a la materia. Para otra materia se añade Enrollment conservando Student y User. El caso de apoyo pertenece al estudiante y periodo; no requiere necesariamente una predicción.")

    page(d, "4 Autorización transacciones y concurrencia")
    h(d, "Identidad y alcance", 2)
    p(d, "role_for y _require validan el actor en los servicios. visible_groups y teacher_enrollments restringen al docente a sus grupos; visible_students, visible_requests y visible_cases aplican el alcance del tutor por estudiante y periodo. El estudiante accede a sus propios registros. No confiar sólo en que un botón o un selector esté oculto.")
    p(d, "El administrador organiza cuentas y relaciones; eso no concede automáticamente las consultas docentes o las notas privadas del tutor. Al reasignar un tutor se actualizan los casos del periodo y sus versiones, con auditoría. El tutor anterior pierde su acceso por el alcance de las consultas.")
    h(d, "Sesión y credenciales Cloud", 2)
    p(d, "cloud/auth.py autentica con Django y guarda ID, hash de sesión, inicio y rol. get_session_user comprueba cuenta activa, rol y hash en cada ejecución; la duración máxima es de ocho horas. logout borra todo el estado de sesión, incluidos formularios y archivos temporales de accesos.")
    p(d, "Los intentos fallidos se registran en MySQL por cuenta normalizada y comparten el bloqueo entre sesiones. Streamlit no usa la IP declarada por el navegador como garantía de identidad. El portal Django tiene su propia capa de sesión y autenticación; no sustituirla por la sesión Cloud.")
    h(d, "Control de cambios", 2)
    p(d, "save_indicator, import_indicators y add_followup reciben expected_version. Las escrituras bloquean los registros pertinentes y comparan la versión actual con la del formulario; una diferencia exige recargar. Las importaciones académicas se guardan de forma atómica para no sobrescribir sólo una parte del grupo.")
    p(d, "Las altas de organización usan transaction.atomic, full_clean y auditoría. enroll_student bloquea Student y CourseGroup; la importación de inscripciones bloquea los mismos padres en un orden estable. Las restricciones de la base siguen siendo la defensa frente a altas concurrentes. SQLite de pruebas no acredita el comportamiento de select_for_update en MySQL.")
    h(d, "Configuración de producción", 2)
    p(d, "Cloud exige MySQL externo, certificado CA, verificación TLS, secretos privados, hosts exactos y orígenes HTTPS. La aplicación se detiene si falta configuración, conexión o migraciones; no cambia silenciosamente a SQLite. Mantener CORS y XSRF de Streamlit activos y evitar secretos en Git, logs o ejemplos.")

    page(d, "5 Implementación de la inscripción por archivo")
    h(d, "Entrada y revisión", 2)
    p(d, "cloud/enrollment_files.parse(content, filename) acepta .xlsx y CSV UTF-8. Exige las columnas codigo, nombres, apellidos y usuario sin duplicar; admite cambio de orden. Usa Alumnos si existe o la hoja activa. Rechaza fórmulas, columnas desconocidas, datos fuera de columnas, entradas vacías y límites excedidos; omite filas totalmente vacías.")
    p(d, "El límite es 1000 alumnos, 5 MB y 20 MB de XLSX expandido. CSV admite coma, punto y coma o tabulador. workbook genera la plantilla y la hoja Instrucciones; para accesos añade clave y fuerza texto en las celdas exportadas, evitando que un dato se interprete como fórmula.")
    p(d, "prepare(actor, records, group_ids) exige administrador y entre 1 y 30 grupos con periodo abierto. Reserva códigos explícitos y existentes antes de generar EST- automáticamente; el usuario por defecto es alumno. seguido del código en minúsculas. Los alumnos nuevos requieren nombres y apellidos. Los existentes se reconocen por código y deben estar activos con rol estudiante; los campos adicionales deben coincidir, sin actualizar la cuenta.")
    h(d, "Contrato de la vista previa", 2)
    table(d, ["Clave", "Contenido"], [
        ["rows", "codigo, nombres, apellidos, usuario y student_id existente o None."],
        ["group_ids", "Identificadores normalizados y ordenados de materias destino."],
        ["new_accounts", "Cantidad de alumnos que requieren cuenta nueva."],
        ["new_enrollments y skipped", "Relaciones por crear y relaciones que ya existen."],
    ], [2.1, 4.65])
    h(d, "Confirmación y persistencia", 2)
    p(d, "commit vuelve a autorizar, bloquea estudiantes y grupos y repite prepare dentro de transaction.atomic. Compara el plan con la revisión; si cambió, aborta. Cada cuenta nueva usa secrets.token_urlsafe(18), create_account y el hash de Django. Las inscripciones y sus AuditEvent se insertan por lotes de 500; una colisión o fallo revierte cuentas, perfiles, estudiantes, relaciones y auditorías del lote.")
    p(d, "La interfaz vincula la revisión al SHA-256 del contenido y a las materias seleccionadas. Cambiar cualquiera invalida la revisión. Los accesos nuevos quedan temporalmente en la sesión para descargarlos; no se incluyen en la vista previa ni se guardan en texto plano en MySQL. La huella de un lote guardado evita repetirlo en esa sesión, pero no deduplica personas sin código entre sesiones.")
    p(d, "Pruebas: cloud/tests/test_enrollment_import.py cubre permisos, revisión sin escritura, altas, varias materias, alumnos existentes, contraseñas, colisiones, rollback, planes obsoletos, formatos, límites y el recorrido de Streamlit AppTest.")

    page(d, "6 Algoritmo y modelo demostrativo")
    p(d, "La implementación utiliza StandardScaler y MLPClassifier con tres entradas: nota_parcial, asistencia y tareas_entregadas. El ID, nombre y resultado_final no entran en la inferencia. El objetivo sintético es binario: 0 aprobado y 1 reprobado. La MLP se conserva por el alcance didáctico del proyecto.")
    table(d, ["Parámetro", "Valor del experimento"], [
        ["hidden_layer_sizes", "(8,), una capa oculta de ocho neuronas"],
        ["activation y solver", "relu y adam, valores predeterminados de la versión fijada"],
        ["max_iter", "1000"],
        ["early_stopping y validation_fraction", "True y 0.15, reserva interna del entrenamiento"],
        ["n_iter_no_change y random_state", "30 y 2026"],
        ["learning_rate_init y alpha", "0.001 y 0.0001, valores predeterminados"],
        ["Umbral", "0.5, fijo para la demostración"],
    ], [3.0, 3.75])
    p(d, "scripts/train_demo.py separa 240 filas en 144 de entrenamiento, 48 de validación externa y 48 de prueba reservada mediante estratificación y semilla 2026. El escalador aprende sólo de entrenamiento. La reserva interna de early_stopping no sustituye la validación externa ni la prueba. También compara una regla nota < 6 y una regresión logística.")
    meta = json.loads((ROOT / "models/messi_demo.json").read_text(encoding="utf-8"))
    holdout = meta["evaluations"]["test_holdout"]
    table(d, ["Candidato", "Exactitud", "Recall clase 1", "ROC AUC"], [
        [name, f"{values['accuracy']:.2%}", f"{values['recall_risk']:.2%}", f"{values['roc_auc']:.3f}"] for name, values in holdout.items()
    ], [2.65, 1.2, 1.5, 1.4])
    p(d, "Los valores anteriores proceden de models/messi_demo.json del experimento sintético local. No prueban superioridad de la MLP, eficacia educativa, equidad ni calibración para datos reales. La interfaz Cloud presenta score multiplicado por 100; sigue siendo una puntuación de demostración.")
    p(d, "messi/model.py verifica ruta dentro de models, extensión, tamaño, metadatos y SHA-256 antes de abrir el joblib. El hash detecta diferencias respecto al JSON, pero no autentica al autor si se alteran ambos. No aceptar artefactos binarios subidos por visitantes. Ajustar datos, arquitectura o umbral exige reentrenar y verificar metadatos.")

    page(d, "7 Formatos y origen de los datos")
    table(d, ["Uso", "Datos y origen"], [
        ["Entrenamiento", "data/synthetic/training_synthetic.csv. Generador del equipo, semilla 2026; 240 registros artificiales con resultado_final."],
        ["Ejemplo docente", "students_demo.csv y Plantilla_MESSI.xlsx del repositorio. Datos ficticios, sin etiqueta final para inferencia."],
        ["Inscripción Cloud", "codigo, nombres, apellidos, usuario. No incluye indicadores. La lista selecciona materias desde la interfaz."],
        ["Indicadores docentes", "id_estudiante, nota_parcial, asistencia, tareas_entregadas; códigos de alumnos previamente inscritos."],
        ["Persistencia Cloud", "User, Student, Enrollment, Indicator y expedientes relacionados en MySQL mediante Django."],
        ["Persistencia Windows", "messi.sqlite3 del usuario local; su esquema y sus migraciones son independientes."],
    ], [1.8, 4.95])
    p(d, "Los indicadores admiten nota de 0 a 10 y porcentajes de 0 a 100. El contrato general de messi/data.py rechaza valores no finitos, códigos repetidos, columnas extra y fórmulas. La captura por conteos requiere totales enteros positivos y realizados no mayores que el total.")
    h(d, "Plantilla docente con alumnos autorizados", 2)
    p(d, "cloud/import_templates.group_rows obtiene sólo las inscripciones de la materia autorizada. El Excel presenta Estudiante, ID estudiante, Nota primer parcial, Asistencia (%) y Tareas entregadas (%). El nombre es una referencia visual; el código sigue determinando la identidad.")
    p(d, "_teacher_import usa load_excel con allow_temporary_ids=False, allow_student_names=True y skip_blank_indicators=True. La variante CSV omite filas con los tres indicadores vacíos. Las filas parciales se rechazan y cero es válido. Después valida pertenencia al grupo y obtiene expected_versions antes de confirmar con services.import_indicators.")
    p(d, "La importación académica admite hasta 10000 registros y 5 MB, con 50 MB de Excel expandido según messi/data.py. Estos límites son distintos de los 1000 alumnos y 20 MB expandidos del archivo de inscripción. Un CSV exportado no constituye un respaldo del expediente completo.")
    h(d, "Evolución del contrato", 2)
    p(d, "Para añadir un campo de inferencia, ajustar FEATURES, validadores, lectura y exportación, formularios, plantillas, servicios, esquema mediante migraciones y metadatos del modelo. Añadir pruebas de rango y compatibilidad. No incluir nombres o identificadores como entradas del modelo para completar un cambio de presentación.")

    page(d, "8 Preparar desarrollo y configurar Cloud")
    p(d, "Usa Python 3.13 para reproducir el despliegue Cloud. Las dependencias Cloud se fijan en cloud/requirements.txt; requirements-web.txt sirve al portal Django autohospedado; requirements.txt y requirements-build.txt corresponden a escritorio y empaquetado. Ejecuta desde la raíz del repositorio.")
    code(d, "py -3.13 -m venv .venv-web\n.\\.venv-web\\Scripts\\python.exe -m pip install -r cloud/requirements.txt")
    p(d, "Copia los marcadores de cloud/secrets.example.toml a un archivo privado .streamlit/secrets.toml. No copiar credenciales de una demostración a la documentación. Cloud lee la sección messi_web y configura los siguientes valores:")
    table(d, ["Clave TOML", "Finalidad"], [
        ["secret_key", "Clave privada Django de al menos 50 caracteres en producción."],
        ["db_host, db_port y db_name", "Servidor MySQL externo, puerto y base destinada al portal."],
        ["db_user y db_password", "Cuenta de ejecución con los permisos necesarios para los datos."],
        ["allowed_hosts", "Nombres de host explícitos, sin protocolo ni comodines."],
        ["csrf_trusted_origins", "Orígenes HTTPS exactos, sin ruta ni comodines."],
        ["ssl_ca_pem", "Certificado CA del proveedor, sin clave privada."],
    ], [2.7, 4.05])
    p(d, "La interfaz Cloud no crea tablas o administradores al abrirse. Inicializa explícitamente la base desde un entorno administrativo y luego usa credenciales de ejecución adecuadas:")
    code(d, ".\\.venv-web\\Scripts\\python.exe cloud\\runtime.py migrate\n.\\.venv-web\\Scripts\\python.exe cloud\\runtime.py createsuperuser\n.\\.venv-web\\Scripts\\python.exe cloud\\runtime.py check")
    p(d, "Puedes indicar --secrets con una ruta privada alternativa. El comando prepara el entorno sin sobrescribir web/.env. Este último archivo es la configuración independiente del portal Django autohospedado; consulta docs/USO_WEB.md para HTTPS, Docker y servidor WSGI.")
    p(d, "Para cambiar una función, ubica la regla en un servicio y usa la interfaz para recoger y presentar datos. Preserva autorización, transacción, auditoría y versiones. Una actualización de campos persistentes debe incluir migración; no borrar tablas para simplificar una actualización.")

    page(d, "9 Verificar publicar y mantener")
    h(d, "Pruebas en una base aislada", 2)
    p(d, "Nunca ejecutar las suites de escritura contra la base de producción. Para comprobaciones funcionales aisladas puede habilitarse SQLite de forma deliberada en desarrollo:")
    code(d, "$env:MESSI_WEB_ENV = 'development'\n$env:MESSI_WEB_TEST_SQLITE = '1'\n$env:MESSI_WEB_SQLITE_PATH = ':memory:'\n$env:PYTHONPATH = (Get-Location).Path\n.\\.venv-web\\Scripts\\python.exe web\\manage.py test portal.tests cloud.tests --noinput\nRemove-Item Env:\\MESSI_WEB_TEST_SQLITE\nRemove-Item Env:\\MESSI_WEB_SQLITE_PATH")
    p(d, "Ejecuta además la misma suite con MySQL de ensayo y una cuenta autorizada para crear la base test_*. Las pruebas MySQL comprueban bloqueo y concurrencia; las omitidas no cuentan como aprobadas. .github/workflows/messi-web.yml ejecuta migraciones, verificación del esquema, portal.tests, cloud.tests y construcción Docker con MySQL 8.4. pruebas.yml mantiene la suite histórica de escritorio y su integración MySQL.")
    p(d, "La carga por archivo se verificó con 41 pruebas de importación, organización y plantillas; el cambio publicado también pasó las comprobaciones GitHub del portal, integración MySQL y escritorio. Son resultados de esa revisión, no una garantía permanente ni un ensayo con alumnos reales.")
    h(d, "Publicar en Community Cloud", 2)
    steps(d, ["Sube el cambio revisado a GitHub y comprueba sus pruebas. El portal publicado sigue main y usa cloud/app.py.", "En Community Cloud configura Python 3.13, cloud/requirements.txt, packages.txt y Secrets privados. Mantén activas las protecciones de Streamlit.", "Si cambia el esquema, aplica las migraciones explícitamente con la configuración y cuenta adecuadas antes de habilitar el código que las requiere. Haz respaldo y planifica compatibilidad.", "Después de publicar o cambiar Secrets, espera el reinicio y verifica inicio de sesión, roles, grupos, revisión de archivos y persistencia desde sesiones independientes."])
    p(d, "Para una regresión de código, conserva la revisión Git anterior y prepara una reversión. Revertir código no revierte datos ni migraciones. El restablecimiento de una base debe ensayarse aparte y no sustituirse por un git reset.")
    h(d, "Respaldo y soporte", 2)
    p(d, "El proveedor MySQL externo debe tener respaldos y una restauración comprobada en otra base. scripts/respaldo_web.py ofrece backup y restore para el despliegue Docker o MySQL configurado en web/.env con --local; no toma automáticamente los Secrets de Aiven ni replica su verificación TLS. Usa la guía del proveedor para esa instancia y protege los archivos de respaldo.")

    page(d, "10 Extender el escritorio y conservar sus datos")
    p(d, "El escritorio conserva su construcción independiente. Para desarrollarlo instala requirements-build.txt en .venv, ejecuta scripts/generate_synthetic.py y scripts/train_demo.py --seed 2026 si necesitas regenerar el modelo, y verifica tests con unittest. Para el portal Cloud el modelo puede generarse durante la inicialización cuando falta.")
    code(d, ".\\.venv\\Scripts\\python.exe -m unittest discover -s tests -v\npowershell -ExecutionPolicy Bypass -File scripts\\build_windows.ps1")
    p(d, "La construcción usa PyInstaller e Inno Setup 6. -SoloPortable produce dist/MESSI; se debe distribuir la carpeta completa con sus recursos. El instalador queda en release/MESSI-Setup-Windows-x64.exe. Comprueba --self-test y --smoke-server desde otra carpeta antes de entregar.")
    p(d, "Los recursos instalados se separan de %LOCALAPPDATA%\\MESSI\\data\\messi.sqlite3. Respaldar mediante la API SQLite contempla el modo WAL. Una restauración exige cerrar la aplicación, conservar la base destino y apartar sus archivos -wal y -shm antes de colocar el respaldo. Esta base no sincroniza equipos ni se transforma automáticamente en la base Django.")
    page(d, "11 Créditos dependencias y licencias")
    p(d, "El equipo MESSI conserva la autoría del proyecto. La documentación y la integración se prepararon con asistencia de Codex; los datos de entrenamiento y sus etiquetas fueron generados artificialmente. No se atribuyen pruebas humanas sin su evidencia. Una licencia del código propio debe establecerse expresamente por el equipo.")
    table(d, ["Dependencia", "Cloud fijada", "Licencia de referencia"], [
        ["Streamlit", "1.65.0", "Apache 2.0"], ["Django", "5.2.18", "BSD de tres cláusulas"],
        ["mysqlclient", "2.3.0", "GPL versión 2 o posterior"], ["WhiteNoise", "6.12.0", "MIT"],
        ["pandas y scikit-learn", "2.3.3 y 1.7.2", "BSD de tres cláusulas"],
        ["NumPy y SciPy", "2.3.3 y 1.16.2", "BSD de tres cláusulas"],
        ["joblib y openpyxl", "1.5.2 y 3.1.5", "BSD de tres cláusulas y MIT"],
        ["python-dotenv y tzdata", "1.1.1 y 2026.2", "BSD de tres cláusulas y Apache 2.0; conservar avisos de IANA"],
    ], [2.6, 1.3, 2.85])
    p(d, "Conserva los avisos LICENSE y metadatos de cada distribución y revisa las condiciones al redistribuir o actualizar. Python usa la licencia PSF y SQLite es de dominio público. El empaquetado añade las licencias de PyInstaller, Inno Setup y Tcl/Tk; no asumir que la licencia del portal cubre todos los componentes del instalador.")
    h(d, "Referencias de implementación", 2)
    p(d, "Código y guías: README.md; docs/STREAMLIT_CLOUD.md; docs/USO_WEB.md; docs/FLUJO_CONFIGURACION_ESCUELA.md; cloud/requirements.txt; models/messi_demo.json. Repositorio: https://github.com/a3524110303-cpu/M.E.S.S.I-")
    p(d, "Documentación oficial consultada: Django 5.2, https://docs.djangoproject.com/en/5.2/ ; Streamlit, https://docs.streamlit.io/deploy/streamlit-community-cloud/get-started/trust-and-security ; MLP de scikit-learn, https://scikit-learn.org/1.7/modules/neural_networks_supervised.html . Las versiones instaladas se determinan por los archivos fijados del proyecto.")
    d.save(target)
    return target


if __name__ == "__main__":
    for output in (user_manual(), programmer_manual()):
        print(output)
